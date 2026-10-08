import openpyxl

FILE = r"C:\Users\ASUS\Documents\Manufacturing\3Marketing\Recap\files_BT\Zales Disney Bridal recap 29.09.xlsx"

wb = openpyxl.load_workbook(FILE)
print("=== SHEETS ===")
print(wb.sheetnames)

ws = wb.active
print(f"\n=== ACTIVE SHEET: '{ws.title}' | Max Row: {ws.max_row} | Max Col: {ws.max_column} ===\n")

print("=== ALL ROWS (non-empty) ===")
for r in range(1, ws.max_row + 1):
    row_data = {}
    for c in range(1, ws.max_column + 1):
        val = ws.cell(r, c).value
        if val is not None:
            row_data[c] = val
    if row_data:
        print(f"Row {r:>3}: {row_data}")

print("\n=== MERGED CELLS ===")
for m in ws.merged_cells.ranges:
    print(m)
