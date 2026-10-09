"""
recap_engine.builder
====================
Updates and formats Client Recap presentation workbooks using openpyxl.
Maintains 100% fidelity of cell fonts, colors, borders, alignments, dynamic row sizing,
and complex merged cells across headers and terms-and-conditions footers.
"""

import os
import re
import copy
from datetime import date
from typing import List, Dict, Tuple, Any, Optional
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill, Border

from .models import StyleItem, RecapRow
from .resolvers import QualityCodeResolver, ToleranceResolver, DescriptionHelper


class RecapWorkbookBuilder:
    """
    Builder responsible for generating and formatting the updated Client Recap Excel workbook.
    """

    DATA_START_ROW = 15

    # Column mappings (Table strictly spans Columns 1 to 11, A to K)
    COL_SR = 1
    COL_STYLE = 2
    COL_IMAGE = 3
    COL_DESC = 4
    COL_METAL = 5
    COL_METAL_WT = 6
    COL_DIA_QLY = 7
    COL_CTTW = 8
    COL_GEM_INFO = 9
    COL_PRICE = 10
    COL_COMMENTS = 11
    TOTAL_COLUMNS = 11

    @staticmethod
    def load_template_metadata(template_path: str) -> Tuple[Dict[str, str], Dict[str, str], Dict[str, str]]:
        """
        Reads existing metadata from the recap template (CTTW, descriptions, dia quality overrides).
        """
        if os.path.basename(template_path).lower() == "recap-template.xlsx":
            return {}, {}, {}

        wb = openpyxl.load_workbook(template_path, data_only=True)
        ws = wb.active
        cttw_map: Dict[str, str] = {}
        desc_map: Dict[str, str] = {}
        dia_map: Dict[str, str] = {}

        for r in range(15, ws.max_row + 1):
            st = ws.cell(r, 2).value
            ct = ws.cell(r, 8).value
            ds = ws.cell(r, 4).value
            dq = ws.cell(r, 7).value
            if st:
                m = re.search(r'\d{6}', str(st))
                if m:
                    dig = m.group(0)
                    if ct: cttw_map[dig] = str(ct).strip()
                    if ds: desc_map[dig] = str(ds).strip()
                    if dq: dia_map[dig]  = str(dq).strip()

        wb.close()
        return cttw_map, desc_map, dia_map

    @classmethod
    def prepare_recap_rows(
        cls,
        items: List[StyleItem],
        code_resolver: QualityCodeResolver,
        tolerance_resolver: ToleranceResolver,
        template_metadata: Tuple[Dict[str, str], Dict[str, str], Dict[str, str]],
        collection_name: str = "",
        price_basis: str = "MEMO_COST"
    ) -> List[RecapRow]:
        """
        Transforms parsed StyleItem models into presentation RecapRow objects.
        """
        recap_cttw, recap_desc, recap_dia = template_metadata
        rows: List[RecapRow] = []

        for sr, item in enumerate(items, start=1):
            m = re.search(r'\d{6}', item.style_no)
            num = m.group(0) if m else item.style_no

            # 1. Carat weight fraction (Tolerance Resolution strictly on diamonds)
            has_dia = item.has_diamonds
            tot_wt = item.total_diamond_weight if has_dia else 0.0
            frac = tolerance_resolver.get_fraction(tot_wt) if (has_dia and tot_wt > 0) else ""

            if num in recap_cttw:
                cttw = recap_cttw[num]
            elif not has_dia:
                cttw = "-"
            else:
                center_note = ""
                if item.dia1.count == 1 and item.dia1.carat_weight > 0:
                    center_info = DescriptionHelper.get_center_stone_desc(item.dia1.shape_size, item.dia1.carat_weight)
                    if center_info:
                        center_note = f"      ({center_info})"
                cttw = f"{frac} cttw{center_note}" if frac else (f"{tot_wt:.2f} cttw{center_note}" if tot_wt > 0 else "-")

            # 2. Quality descriptions
            if has_dia:
                dia_qly_desc = code_resolver.resolve_diamond_quality(item.dia1.quality_code)
            else:
                dia_qly_desc = "-"

            # Gemstone Information (Column I / 9):
            gem_qly_desc = ""
            raw_gem_code = (item.gem.quality_code or "").strip().upper()
            if raw_gem_code and raw_gem_code in code_resolver.stone_map:
                base_gem_name = code_resolver.stone_map[raw_gem_code]
            else:
                base_gem_name = item.gem.label if item.gem.label and item.gem.label != 'Gem' else ""

            clean_gem_name = ""
            if 'MORGANITE' in str(base_gem_name).upper() or 'MORGANITE' in str(item.gem.dimension_mm).upper():
                clean_gem_name = "Morganite"
            elif 'SAPPHIRE' in str(base_gem_name).upper() or 'SAPH' in str(base_gem_name).upper():
                clean_gem_name = "Cr.White Sapphire"
            elif base_gem_name:
                clean_gem_name = base_gem_name

            if base_gem_name or item.gem.shape_size or item.gem.dimension_mm:
                parts_g = []
                if item.is_bridal_set or 'MORGANITE' in str(base_gem_name).upper() or 'MORGANITE' in str(item.gem.dimension_mm).upper():
                    parts_g.append("CENTER MORGANITE")
                elif base_gem_name:
                    parts_g.append(base_gem_name)

                dim_str = item.gem.dimension_mm.replace(' X ', 'x').strip() if item.gem.dimension_mm else ""
                sh_str = item.gem.shape_size.strip() if item.gem.shape_size else ""

                if dim_str and dim_str.upper() not in [p.upper() for p in parts_g]:
                    parts_g.append(dim_str)
                elif sh_str and sh_str.upper() not in [p.upper() for p in parts_g]:
                    parts_g.append(sh_str)

                gem_qly_desc = " ".join(parts_g).strip()
            else:
                gem_qly_desc = "-" if has_dia else ""

            # 3. Description (uses commercial fraction e.g. '5/8 ctw' instead of raw cost or decimals)
            if num in recap_desc:
                description = recap_desc[num]
            else:
                clean_frac = frac.strip() if frac and not re.search(r'^\d+\.\d+$', frac.strip()) else ""
                ctw_str = f"{clean_frac}cttw" if clean_frac else ""
                description = DescriptionHelper.build_item_description(
                    metal_desc=item.metal.description_text,
                    ctw_desc=ctw_str,
                    dia_desc=dia_qly_desc if has_dia else "",
                    gem_desc=gem_qly_desc if gem_qly_desc != '-' else "",
                    style_no=item.style_no,
                    character_name=item.character_name,
                    collection_name=collection_name,
                    has_diamonds=has_dia,
                    gem_name=clean_gem_name
                )

            final_dia_qly = recap_dia.get(num, dia_qly_desc)
            if final_dia_qly == "NA":
                final_dia_qly = "-"
            if cttw == "NA":
                cttw = "-"

            # 4. Pricing based on selected basis
            basis_clean = (price_basis or "MEMO_COST").upper()
            if "SELLING" in basis_clean or "COOP" in basis_clean:
                if item.is_bridal_set and item.box_price is not None:
                    unit_price = round(item.box_price, 2)
                elif item.ny_sell_coop is not None:
                    unit_price = round(item.ny_sell_coop, 2)
                else:
                    unit_price = item.effective_unit_price
            else:
                unit_price = item.effective_unit_price

            rows.append(RecapRow(
                sr=sr,
                style_no=item.style_no,
                image="",
                description=description,
                metal=item.metal.cell_text,
                metal_wt=item.metal.display_weight,
                dia_qly=final_dia_qly,
                cttw=cttw,
                gem_info=gem_qly_desc,
                unit_price=unit_price,
                comments=""
            ))

        return rows

    @staticmethod
    def _apply_cell_style(cell: openpyxl.cell.Cell, style_dict: Dict[str, Any]) -> None:
        if style_dict.get('font'):          cell.font          = style_dict['font']
        if style_dict.get('fill'):          cell.fill          = style_dict['fill']
        if style_dict.get('border'):        cell.border        = style_dict['border']
        if style_dict.get('alignment'):     cell.alignment     = style_dict['alignment']
        if style_dict.get('number_format'): cell.number_format = style_dict['number_format']

    @staticmethod
    def _ensure_merged(ws: openpyxl.worksheet.worksheet.Worksheet, coord: str) -> None:
        existing = {m.coord for m in ws.merged_cells.ranges}
        if coord not in existing:
            try:
                ws.merge_cells(coord)
            except Exception:
                pass

    def update_recap(
        self,
        template_path: str,
        rows: List[RecapRow],
        output_path: str,
        customer_name: Optional[str] = None,
        gold_lock_rate: float = 4250.0,
        silver_lock_rate: float = 65.0
    ) -> str:
        """
        Populates rows into the client recap template and saves the updated workbook.
        """
        wb = openpyxl.load_workbook(template_path)
        try:
            ws = wb.active
            table_cols = self.TOTAL_COLUMNS

            # 1. Update quotation date at (Row 6, Col 7)
            ws.cell(row=6, column=7).value = date.today().strftime("%B %d, %Y")
            if customer_name:
                ws.cell(row=9, column=1).value = customer_name

            # Metal Lock rate prices at (Row 12, Col 1)
            ws.cell(row=12, column=1).value = f"  Gold Lock: ${gold_lock_rate:,.2f} / troy oz     |     Silver Lock: ${silver_lock_rate:,.2f} / troy oz"

            # 2. Locate dynamic TERMS & CONDITIONS row
            tc_row = None
            for r in range(self.DATA_START_ROW, ws.max_row + 1):
                val = ws.cell(row=r, column=1).value
                if val and 'TERMS' in str(val).upper():
                    tc_row = r
                    break
            assert tc_row, "Could not locate TERMS & CONDITIONS row in template!"

            # 3. Capture baseline styles and height from single template row (Row 15)
            base_row = self.DATA_START_ROW
            base_height = ws.row_dimensions[base_row].height or 62.25

            row_styles_base: Dict[int, Dict[str, Any]] = {}
            for c in range(1, table_cols + 1):
                cell = ws.cell(row=base_row, column=c)
                row_styles_base[c] = {
                    'font': copy.copy(cell.font) if cell.has_style and cell.font else None,
                    'fill': copy.copy(cell.fill) if cell.has_style and cell.fill else None,
                    'border': copy.copy(cell.border) if cell.has_style and cell.border else None,
                    'alignment': copy.copy(cell.alignment) if cell.has_style and cell.alignment else None,
                    'number_format': cell.number_format,
                }

            # Setup alternating zebra striping styles:
            # alt 0 (even rows): original template row fill (soft off-white FFFAFAFA / FFFFFFFF)
            # alt 1 (odd rows): subtle light tint (FFF4F5F8)
            zebra_fill = PatternFill(fill_type="solid", start_color="FFF4F5F8", end_color="FFF4F5F8")

            template_styles: Dict[int, Dict[int, Dict[str, Any]]] = {0: {}, 1: {}}
            for c in range(1, table_cols + 1):
                template_styles[0][c] = copy.copy(row_styles_base[c])
                odd_style = copy.copy(row_styles_base[c])
                odd_style['fill'] = zebra_fill
                template_styles[1][c] = odd_style

            # 4. Calculate dynamic row adjustment
            new_count = len(rows)
            target_tc_row = self.DATA_START_ROW + new_count
            diff = target_tc_row - tc_row

            # 5. Shift footer merges cleanly
            footer_merges = []
            if diff != 0:
                for m in list(ws.merged_cells.ranges):
                    if m.min_row >= tc_row:
                        footer_merges.append((m.min_col, m.min_row, m.max_col, m.max_row, m.coord))
                        ws.unmerge_cells(m.coord)
                    elif diff < 0 and m.min_row >= target_tc_row:
                        ws.unmerge_cells(m.coord)

                if diff < 0:
                    ws.delete_rows(target_tc_row, -diff)
                elif diff > 0:
                    ws.insert_rows(tc_row, diff)

                for (c1, r1, c2, r2, _) in footer_merges:
                    new_coord = f"{get_column_letter(c1)}{r1 + diff}:{get_column_letter(c2)}{r2 + diff}"
                    ws.merge_cells(new_coord)

            # 6. Clear placeholder values
            for r in range(self.DATA_START_ROW, self.DATA_START_ROW + new_count):
                for c in range(1, table_cols + 1):
                    ws.cell(row=r, column=c).value = None

            # 7. Write data rows applying captured styles and row heights
            for i, row_data in enumerate(rows):
                row_num = self.DATA_START_ROW + i
                alt = i % 2

                ws.row_dimensions[row_num].height = base_height

                for c in range(1, table_cols + 1):
                    self._apply_cell_style(ws.cell(row=row_num, column=c), template_styles[alt][c])

                # Clean any accidental fill/border in columns beyond the table (Column L, M, etc.)
                for c in range(table_cols + 1, ws.max_column + 1):
                    extra_cell = ws.cell(row=row_num, column=c)
                    extra_cell.fill = PatternFill(fill_type=None)
                    extra_cell.border = Border()

                ws.cell(row=row_num, column=self.COL_SR).value = row_data.sr
                ws.cell(row=row_num, column=self.COL_STYLE).value = row_data.style_no
                ws.cell(row=row_num, column=self.COL_IMAGE).value = row_data.image
                ws.cell(row=row_num, column=self.COL_DESC).value = row_data.description
                ws.cell(row=row_num, column=self.COL_METAL).value = row_data.metal

                # Metal weight
                m_cell = ws.cell(row=row_num, column=self.COL_METAL_WT)
                m_cell.value = row_data.metal_wt
                if isinstance(row_data.metal_wt, (int, float)):
                    m_cell.number_format = '0.00'

                ws.cell(row=row_num, column=self.COL_DIA_QLY).value = row_data.dia_qly
                ws.cell(row=row_num, column=self.COL_CTTW).value = row_data.cttw
                ws.cell(row=row_num, column=self.COL_GEM_INFO).value = row_data.gem_info
                ws.cell(row=row_num, column=self.COL_COMMENTS).value = row_data.comments

                # Unit price
                price_cell = ws.cell(row=row_num, column=self.COL_PRICE)
                price_cell.value = row_data.unit_price
                if isinstance(row_data.unit_price, (int, float)):
                    price_cell.number_format = '[$$-409]#,##0.00'

            # 8. Restore footer heights and ensure merged cells
            cur_tc = target_tc_row
            ws.row_dimensions[cur_tc].height = 42.0       # TERMS & CONDITIONS header
            ws.row_dimensions[cur_tc + 1].height = 30.0   # T&C validity & terms text
            ws.row_dimensions[cur_tc + 2].height = 15.0   # Spacer row
            ws.row_dimensions[cur_tc + 3].height = 20.0   # Authorized By
            ws.row_dimensions[cur_tc + 5].height = 7.5    # Bottom margin spacer

            # Ensure essential header merges
            for h_merge in [
                "A5:F5", "A8:C8", "G8:K8", "A9:F9", "G9:K9",
                "A11:G11", "H11:K11", "A12:G12", "H12:K12"
            ]:
                self._ensure_merged(ws, h_merge)

            # Ensure footer merges
            self._ensure_merged(ws, f"A{cur_tc}:K{cur_tc}")
            self._ensure_merged(ws, f"A{cur_tc + 1}:K{cur_tc + 1}")
            self._ensure_merged(ws, f"A{cur_tc + 3}:F{cur_tc + 3}")
            self._ensure_merged(ws, f"H{cur_tc + 3}:K{cur_tc + 3}")

            # 9. Save workbook safely
            os.makedirs(os.path.dirname(output_path), exist_ok=True)
            try:
                wb.save(output_path)
                return output_path
            except PermissionError:
                alt_path = output_path.replace(".xlsx", "_new.xlsx")
                wb.save(alt_path)
                return alt_path
        finally:
            wb.close()
