import openpyxl

FILE = r"C:\Users\ASUS\Documents\Manufacturing\3Marketing\Recap\files_BT\Code_Mapping.xlsx"
wb = openpyxl.load_workbook(FILE)
print("Sheets:", wb.sheetnames)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n=== Sheet: '{sheet_name}' | Rows: {ws.max_row} | Cols: {ws.max_column} ===")
    for r in range(1, min(ws.max_row + 1, 40)):
        row_data = {}
        for c in range(1, ws.max_column + 1):
            val = ws.cell(r, c).value
            if val is not None:
                row_data[c] = val
        if row_data:
            print(f"  Row {r:>3}: {row_data}")
