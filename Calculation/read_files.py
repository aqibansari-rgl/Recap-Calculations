import csv
from openpyxl import load_workbook
import os

def extract_data_and_images(excel_path, output_csv_path, images_dir):
    print(f"Loading workbook: {excel_path}...")
    try:
        wb = load_workbook(excel_path, data_only=True)
    except Exception as e:
        print(f"Error loading workbook with openpyxl: {e}")
        print("Please ensure openpyxl is installed (`pip install openpyxl`)")
        return
        
    ws = wb.active

    # 1. Extract Data to CSV
    print(f"Extracting data to {output_csv_path}...")
    with open(output_csv_path, mode='w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        for row in ws.iter_rows(values_only=True):
            writer.writerow(row)
    print("Data extracted successfully.")

    # 2. Extract Images
    print(f"Extracting images to {images_dir}...")
    os.makedirs(images_dir, exist_ok=True)
    
    if hasattr(ws, '_images') and ws._images:
        for idx, img in enumerate(ws._images):
            try:
                # Attempt to get row to help identify the style number
                row_idx = "unknown"
                if hasattr(img.anchor, '_from'):
                    row_idx = img.anchor._from.row + 1
                elif hasattr(img.anchor, 'row'):
                    row_idx = img.anchor.row
                
                # Image data and extension
                img_data = img._data()
                
                # Basic check for image type based on magic numbers
                ext = "png"
                if img_data.startswith(b'\xff\xd8'):
                    ext = "jpg"
                
                img_filename = f"image_row_{row_idx}_{idx}.{ext}"
                img_filepath = os.path.join(images_dir, img_filename)
                
                with open(img_filepath, "wb") as f:
                    f.write(img_data)
                print(f"Saved {img_filename}")
            except Exception as e:
                print(f"Failed to extract an image: {e}")
    else:
        print("No images found in the worksheet using openpyxl.")

if __name__ == "__main__":
    excel_file = "C:/Users/ASUS/Documents/Manufacturing/3Marketing/Recap/files_BT/Zales Dinsey Bridal pricing dtd 29.09.xlsx"
    csv_file = "outputs/Zales_Disney/extracted_data.csv"
    img_folder = "outputs/Zales_Disney/style_images"

    if not os.path.exists(excel_file):
        print("Please provide the path to the Excel file.")
        exit()
    
    os.makedirs(os.path.dirname(csv_file), exist_ok=True)
    os.makedirs(os.path.dirname(img_folder), exist_ok=True)
    

    extract_data_and_images(excel_file, csv_file, img_folder)
    print(f"\nDone. The raw data is saved in '{csv_file}' and images in the '{img_folder}' directory.")
    print("Please share the CSV or let me know the output so we can build the structural parser!")
